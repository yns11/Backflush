"""Tests du pivot de la vue synthétique et de son export Excel.

Deux invariants sont vérifiés en priorité, parce qu'ils portent la demande
métier :

* le pivot conserve les zéros — c'est chaque support qui décide de les afficher
  ou non ;
* le classeur écrit un ``0`` **numérique** dans les croisements nuls, de sorte
  qu'une somme Excel fonctionne sans neutraliser les cases vides.
"""

from __future__ import annotations

from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from app.server.domain.dictionary import GRILLES
from app.server.domain.filters import Filtres
from app.server.domain.synthese import croiser
from app.server.services.export_excel import (
    construire_classeur,
    construire_classeur_synthese,
    nom_fichier_synthese,
)

PRODUCTION = [
    {
        "parent_itemid": "PAR-001", "parent_name": "Stator A", "bomid": "BOM-1",
        "annee": 2026, "semaine": 14, "qty_produite": Decimal("100"),
        "valeur_produite": Decimal("5000"),
    },
    {
        "parent_itemid": "PAR-001", "parent_name": "Stator A", "bomid": "BOM-1",
        "annee": 2026, "semaine": 15, "qty_produite": Decimal("50"),
        "valeur_produite": Decimal("2500"),
    },
    {
        "parent_itemid": "PAR-002", "parent_name": "Stator B", "bomid": "BOM-2",
        "annee": 2026, "semaine": 15, "qty_produite": Decimal("300"),
        "valeur_produite": Decimal("18000"),
    },
]

ECARTS = [
    {
        "child_itemid": "CMP-VIS", "child_name": "Vis M4", "coef_bom": Decimal("4.0000"),
        "annee": 2026, "semaine": 14,
        "ecart_equivalent_produit": Decimal("-2.5"), "ecart_valorise": Decimal("-30"),
    },
    {
        "child_itemid": "CMP-VIS", "child_name": "Vis M4", "coef_bom": Decimal("4.0000"),
        "annee": 2026, "semaine": 15,
        "ecart_equivalent_produit": Decimal("0"), "ecart_valorise": Decimal("0"),
    },
    {
        "child_itemid": "CMP-AIM", "child_name": "Aimant", "coef_bom": Decimal("1"),
        "annee": 2026, "semaine": 15,
        "ecart_equivalent_produit": Decimal("12"), "ecart_valorise": Decimal("900"),
    },
]


class TestCroiser:
    def test_les_semaines_sont_uniques_et_triees(self) -> None:
        modele = croiser(PRODUCTION, ECARTS, en_valeur=False)
        assert [semaine.cle for semaine in modele.semaines] == ["2026-14", "2026-15"]
        assert [semaine.libelle for semaine in modele.semaines] == ["S14", "S15"]

    def test_un_parent_par_ligne_avec_cumul_des_semaines(self) -> None:
        modele = croiser(PRODUCTION, ECARTS, en_valeur=False)
        par_reference = {ligne.reference: ligne for ligne in modele.production}
        assert set(par_reference) == {"PAR-001", "PAR-002"}
        assert par_reference["PAR-001"].valeurs == {"2026-14": 100.0, "2026-15": 50.0}
        assert par_reference["PAR-001"].total == 150.0
        assert par_reference["PAR-001"].complement == "BOM-1"

    def test_la_production_est_classee_par_volume_decroissant(self) -> None:
        modele = croiser(PRODUCTION, ECARTS, en_valeur=False)
        assert [ligne.reference for ligne in modele.production] == ["PAR-002", "PAR-001"]

    def test_les_ecarts_sont_classes_sur_la_valeur_absolue(self) -> None:
        modele = croiser(PRODUCTION, ECARTS, en_valeur=False)
        assert [ligne.reference for ligne in modele.ecarts] == ["CMP-AIM", "CMP-VIS"]
        # Le signe reste porté : il distingue non-consommation et surconsommation.
        assert modele.ecarts[1].total == -2.5

    def test_le_pivot_conserve_les_zeros(self) -> None:
        """C'est chaque support qui décide de les masquer, pas le pivot."""
        modele = croiser(PRODUCTION, ECARTS, en_valeur=False)
        vis = next(ligne for ligne in modele.ecarts if ligne.reference == "CMP-VIS")
        assert vis.valeurs["2026-15"] == 0.0

    def test_totaux_de_production_par_semaine_et_sur_la_periode(self) -> None:
        modele = croiser(PRODUCTION, ECARTS, en_valeur=False)
        assert modele.total_production == {"2026-14": 100.0, "2026-15": 350.0}
        assert modele.total_production_periode == 450.0

    def test_part_ecart_rapporte_la_somme_des_valeurs_absolues(self) -> None:
        modele = croiser(PRODUCTION, ECARTS, en_valeur=False)
        # |−2,5| + |12| = 14,5 équivalents produit pour 450 produits.
        assert modele.part_ecart_pct == 14.5 / 450 * 100

    def test_sans_production_la_part_est_nulle_et_non_une_division_par_zero(self) -> None:
        modele = croiser([], ECARTS, en_valeur=False)
        assert modele.part_ecart_pct == 0.0
        assert modele.total_production_periode == 0.0

    def test_la_mesure_valeur_change_les_grandeurs(self) -> None:
        modele = croiser(PRODUCTION, ECARTS, en_valeur=True)
        assert modele.total_production_periode == 25500.0
        aimant = next(ligne for ligne in modele.ecarts if ligne.reference == "CMP-AIM")
        assert aimant.total == 900.0

    def test_le_coefficient_est_formate_sans_zeros_inutiles(self) -> None:
        modele = croiser(PRODUCTION, ECARTS, en_valeur=False)
        complements = {ligne.reference: ligne.complement for ligne in modele.ecarts}
        assert complements["CMP-VIS"] == "4"
        assert complements["CMP-AIM"] == "1"

    def test_les_libelles_manquants_ne_font_pas_echouer_le_pivot(self) -> None:
        modele = croiser(
            [{"parent_itemid": "P", "annee": 2026, "semaine": 1, "qty_produite": None}],
            [{"child_itemid": "C", "annee": 2026, "semaine": 1,
              "ecart_equivalent_produit": None, "coef_bom": None}],
            en_valeur=False,
        )
        assert modele.production[0].designation == "—"
        assert modele.production[0].complement == "—"
        assert modele.ecarts[0].complement == "—"


class TestClasseurSynthese:
    @staticmethod
    def _classeur(en_valeur: bool = False):
        modele = croiser(PRODUCTION, ECARTS, en_valeur=en_valeur)
        filtres = Filtres(
            perimetres=["LIGNE 1"], date_debut=date(2026, 3, 30), date_fin=date(2026, 4, 13)
        )
        flux = construire_classeur_synthese(
            modele, filtres, perimetre="LIGNE 1", en_valeur=en_valeur
        )
        return load_workbook(flux)

    def test_les_deux_onglets_sont_presents(self) -> None:
        classeur = self._classeur()
        assert "Contexte" in classeur.sheetnames
        assert len(classeur.sheetnames) == 2

    def test_les_croisements_nuls_portent_un_zero_numerique(self) -> None:
        """Exigence explicite : le classeur doit rester calculable sans traitement."""
        feuille = self._classeur().worksheets[0]
        lignes = {ligne[0].value: ligne for ligne in feuille.iter_rows()}
        vis = lignes["CMP-VIS"]
        # Colonnes : réf, désignation, coef, S14, S15, total.
        assert vis[3].value == -2.5
        assert vis[4].value == 0
        assert isinstance(vis[4].value, (int, float))

    def test_la_ligne_de_total_somme_la_production(self) -> None:
        feuille = self._classeur().worksheets[0]
        total = next(
            ligne for ligne in feuille.iter_rows() if ligne[0].value == "TOTAL"
        )
        assert total[3].value == 100
        assert total[4].value == 350
        assert total[5].value == 450

    def test_les_sections_encadrent_les_deux_blocs(self) -> None:
        feuille = self._classeur().worksheets[0]
        titres = [
            ligne[0].value for ligne in feuille.iter_rows()
            if isinstance(ligne[0].value, str) and ligne[0].value[:2] in ("1.", "2.")
        ]
        assert titres[0].startswith("1. PRODUCTION")
        assert "équivalent produit" in titres[1]

    def test_en_valeur_les_titres_de_section_portent_l_euro(self) -> None:
        feuille = self._classeur(en_valeur=True).worksheets[0]
        titres = [
            ligne[0].value for ligne in feuille.iter_rows()
            if isinstance(ligne[0].value, str) and ligne[0].value[:2] in ("1.", "2.")
        ]
        assert titres == ["1. PRODUCTION (€)", "2. ÉCART DE PRÉLÈVEMENT (€)"]

    def test_l_entete_porte_une_colonne_par_semaine(self) -> None:
        feuille = self._classeur().worksheets[0]
        entete = [cellule.value for cellule in next(feuille.iter_rows())]
        assert entete == ["Référence", "Désignation", "BOM / Coef.", "S14", "S15", "Total"]

    def test_le_contexte_rappelle_le_perimetre_et_les_filtres(self) -> None:
        feuille = self._classeur()["Contexte"]
        contexte = {ligne[0].value: ligne[1].value for ligne in feuille.iter_rows()}
        assert contexte["Périmètre"] == "LIGNE 1"
        assert contexte["Mesure"] == "Quantité (unités)"
        assert contexte["Périmètres"] == "LIGNE 1"

    def test_les_volets_figent_les_colonnes_d_identification(self) -> None:
        assert self._classeur().worksheets[0].freeze_panes == "D2"


class TestVoletsDesGrilles:
    """Garde-fou sur un piège d'``openpyxl`` en mode ``write_only``.

    La vue de la feuille est sérialisée à l'écriture de la première ligne :
    fixer ``freeze_panes`` ensuite ne lève aucune erreur et ne fait rien. Le
    défaut est invisible à la lecture du code comme à l'exécution des tests
    d'API — seul un aller-retour par le fichier le révèle.
    """

    def test_l_entete_des_grilles_reste_fige(self) -> None:
        grille = GRILLES["composants"]
        lignes = [{colonne.cle: None for colonne in grille.colonnes} for _ in range(3)]
        classeur = load_workbook(construire_classeur(grille, lignes, Filtres()))
        assert classeur.worksheets[0].freeze_panes == "A2"


class TestNomFichierSynthese:
    def test_le_nom_porte_le_perimetre_assaini_et_la_periode(self) -> None:
        nom = nom_fichier_synthese(
            "LIGNE 1 / STATOR A",
            Filtres(date_debut=date(2026, 3, 30), date_fin=date(2026, 4, 13)),
        )
        assert nom.startswith("backflush_synthese_LIGNE-1---STATOR-A_2026-03-30_2026-04-13_")
        assert nom.endswith(".xlsx")

    def test_sans_dates_le_nom_reste_explicite(self) -> None:
        nom = nom_fichier_synthese("L1", Filtres())
        assert "_debut_fin_" in nom
